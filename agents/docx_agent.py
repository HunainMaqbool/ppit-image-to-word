import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Pt, RGBColor

ALIGNMENT_MAP = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justified": WD_ALIGN_PARAGRAPH.JUSTIFY,
}

FONT_SIZE_MAP = {
    "large": 16,
    "medium": 12,
    "small": 10,
}


def _apply_alignment(paragraph, alignment: str) -> None:
    paragraph.alignment = ALIGNMENT_MAP.get(alignment, WD_ALIGN_PARAGRAPH.LEFT)


def _apply_run_formatting(run, block: dict) -> None:
    run.bold = block.get("bold", False)
    run.italic = block.get("italic", False)
    size_hint = block.get("font_size_hint", "medium")
    run.font.size = Pt(FONT_SIZE_MAP.get(size_hint, 12))


def _apply_block(doc: Document, block: dict) -> None:
    block_type = block.get("block_type", "paragraph")
    text = block.get("text", "").strip()
    if not text:
        doc.add_paragraph()
        return

    if block_type == "heading":
        level = int(block.get("level", 1))
        level = max(1, min(3, level))
        doc.add_heading(text, level=level)

    elif block_type == "bullet_list":
        para = doc.add_paragraph(style="List Bullet")
        run = para.add_run(text)
        _apply_run_formatting(run, block)
        _apply_alignment(para, block.get("alignment", "left"))

    elif block_type == "numbered_list":
        para = doc.add_paragraph(style="List Number")
        run = para.add_run(text)
        _apply_run_formatting(run, block)
        _apply_alignment(para, block.get("alignment", "left"))

    elif block_type in ("caption", "footer"):
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.italic = True
        run.font.size = Pt(10)
        _apply_alignment(para, "center" if block_type == "caption" else "left")

    elif block_type in ("table", "table_cell"):
        # Render table/data blocks in monospace so columns stay aligned
        para = doc.add_paragraph()
        run = para.add_run(text)
        run.font.name = "Courier New"
        run.font.size = Pt(9)
        _apply_alignment(para, "left")

    else:
        para = doc.add_paragraph()
        run = para.add_run(text)
        _apply_run_formatting(run, block)
        _apply_alignment(para, block.get("alignment", "left"))


def build_docx(analysis_result: dict, user_overrides: dict | None = None) -> io.BytesIO:
    if user_overrides is None:
        user_overrides = {}

    doc = Document()
    blocks = analysis_result.get("blocks", [])

    for block in blocks:
        block_id = block.get("id", "")
        effective_block = {**block, **user_overrides.get(block_id, {})}
        _apply_block(doc, effective_block)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def build_passthrough_docx(raw_text: str) -> io.BytesIO:
    doc = Document()
    for line in raw_text.splitlines():
        if line.strip():
            doc.add_paragraph(line.strip())
        else:
            doc.add_paragraph()
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
